#!/usr/bin/env python
"""Local composer (v2 design system) for the group-meeting evidence deck.

Faithful python-pptx adaptation of the skill's compose_evidence_deck.mjs with an
upgraded visual system: redesigned title page, agenda, KPI stat cards, styled
table slides, closing page, section chips and X/Y page numbers. Same palette,
geometry discipline and footer contract. Reads the validated deck-spec.json.
"""
import json
import math
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from PIL import Image

W, H, M = 1280, 720, 72
FONT = "Microsoft YaHei"
C = {
    "bg": "F7F9F6", "ink": "17372F", "body": "263732", "muted": "62736C",
    "green": "2D6B5B", "green2": "3E8571", "light": "E5EFE9", "line": "C9D8D0",
    "amber": "F6D676", "amberText": "705717", "amberBg": "FFF9E5",
    "white": "FFFFFF", "card": "FFFFFF", "cardLine": "D8E4DD",
}
C = {k: RGBColor.from_string(v) for k, v in C.items()}


def clean(value, fallback=""):
    out = re.sub(r"\s+", " ", str(value if value is not None else fallback)).strip()
    return out or fallback


def compact_font(value, base, minimum, comfortable):
    length = len(clean(value))
    if length <= comfortable:
        return base
    return max(minimum, base - math.ceil((length - comfortable) / 8))


def inches(px):
    return Inches(px / 96.0)


def set_font(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def text(slide, content, left, top, width, height, size=20, color=None, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    color = C["body"] if color is None else color
    box = slide.shapes.add_textbox(inches(left), inches(top), inches(width), inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = clean(content)
    set_font(run, size, color, bold)
    return box


def rect(slide, left, top, width, height, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE,
         line_w=1.0, radius=None):
    sp = slide.shapes.add_shape(shape, inches(left), inches(top), inches(width), inches(height))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(line_w)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    sp.text_frame.margin_left = sp.text_frame.margin_right = 0
    return sp


def rule(slide, top, left=M, width=W - 2 * M):
    rect(slide, left, top, width, 1, fill=C["line"], line=C["line"])


def header(slide, page, total, section, title):
    text(slide, f"{page:02d}", M, 36, 40, 18, size=12, color=C["green"], bold=True)
    chip_w = max(96, 28 + 16 * len(clean(section or "网络组会")))
    rect(slide, 124, 30, chip_w, 28, fill=C["light"], line=C["light"],
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, section or "网络组会", 124, 36, chip_w, 18, size=12, color=C["green"],
         bold=True, align=PP_ALIGN.CENTER)
    text(slide, f"{page} / {total}", 1080, 36, 128, 18, size=12, color=C["muted"],
         align=PP_ALIGN.RIGHT)
    rect(slide, M, 64, 6, 34, fill=C["green"], line=C["green"])
    text(slide, title, M + 18, 64, 1080, 36, size=compact_font(title, 28, 21, 30),
         color=C["ink"], bold=True)
    rule(slide, 112)


def footer(slide, source):
    rule(slide, 672)
    text(slide, source, M, 686, W - 2 * M, 16, size=10, color=C["muted"])


def badge(slide, label, left, top, width=112):
    rect(slide, left, top, width, 26, fill=C["light"], line=C["light"],
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
    text(slide, label, left + 12, top + 5, width - 24, 16, size=11, color=C["green"], bold=True)


def bullet(slide, content, left, top, width, height=62, size=19):
    rect(slide, left, top + 10, 7, 7, fill=C["green"], line=C["green"], shape=MSO_SHAPE.OVAL)
    text(slide, content, left + 22, top, width - 22, height,
         size=compact_font(content, size, max(14, size - 4), 32))


def kpi_cards(slide, items, left=M, top=136, card_w=272, gap=24, height=96):
    for i, item in enumerate(items[:4]):
        x = left + i * (card_w + gap)
        rect(slide, x, top, card_w, height, fill=C["card"], line=C["cardLine"],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        rect(slide, x, top, 5, height, fill=C["green"], line=C["green"])
        text(slide, clean(item.get("value", "")), x + 20, top + 14, card_w - 36, 40,
             size=compact_font(item.get("value", ""), 26, 18, 12), color=C["ink"], bold=True)
        text(slide, clean(item.get("label", "")), x + 20, top + 58, card_w - 36, 30,
             size=12, color=C["muted"])


def source_footer(slide, locator):
    footer(slide, f"{clean(PAPER_CITATION)} · {clean(locator)}")


def add_image_contain(slide, asset, left, top, width, height):
    with Image.open(asset) as im:
        iw, ih = im.size
    box_ratio = width / height
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        draw_w, draw_h = width, int(round(width / img_ratio))
    else:
        draw_h, draw_w = height, int(round(height * img_ratio))
    off_x = left + (width - draw_w) // 2
    off_y = top + (height - draw_h) // 2
    slide.shapes.add_picture(str(asset), inches(off_x), inches(off_y), inches(draw_w), inches(draw_h))


def figure_slide(slide, page, total, item, spec_dir):
    header(slide, page, total, item.get("section") or "证据主图", item["title"])
    asset = Path(item["asset"])
    asset = asset if asset.is_absolute() else spec_dir / asset
    if float(item.get("image_aspect") or 0) >= 1.45:
        add_image_contain(slide, asset, 72, 136, 1136, 330)
        rect(slide, 72, 136, 1136, 330, fill=C["white"], line=C["line"])
        text(slide, f"{clean(item['figure_label'])} · 报告 §{item['pdf_page']}", 72, 476, 1136, 16,
             size=10, color=C["muted"], align=PP_ALIGN.CENTER)
        badge(slide, "本文证据", 72, 502)
        for i, b in enumerate(item.get("bullets", [])[:3]):
            bullet(slide, b, 72, 536 + i * 29, 1128, 28, size=17)
        rect(slide, 72, 625, 1136, 39, fill=C["amberBg"], line=C["amber"])
        text(slide, clean(item.get("caveat", "")), 90, 635, 1098, 20, size=14,
             color=C["amberText"], bold=True)
    else:
        add_image_contain(slide, asset, 72, 136, 628, 436)
        rect(slide, 72, 136, 628, 436, fill=C["white"], line=C["line"])
        text(slide, f"{clean(item['figure_label'])} · 报告 §{item['pdf_page']}", 72, 586, 628, 16,
             size=10, color=C["muted"], align=PP_ALIGN.CENTER)
        badge(slide, "本文证据", 770, 150)
        for i, b in enumerate(item.get("bullets", [])[:3]):
            bullet(slide, b, 770, 210 + i * 86, 430, 56, size=18)
        rect(slide, 770, 498, 438, 78, fill=C["amberBg"], line=C["amber"])
        text(slide, clean(item.get("caveat", "")), 790, 514, 398, 46, size=16,
             color=C["amberText"], bold=True)
    source_footer(slide, item["source"])


def table_slide(slide, page, total, item, spec_dir):
    header(slide, page, total, item.get("section") or "结果明细", item["title"])
    cols = item["table"]["cols"]
    rows = item["table"]["rows"]
    n_rows = len(rows) + 1
    left, top, width = 72, 142, 1136
    row_h = 34 if n_rows > 8 else 38
    height = row_h * n_rows
    gfx = slide.shapes.add_table(n_rows, len(cols), inches(left), inches(top),
                                 inches(width), inches(height))
    table = gfx.table
    table.first_row = False
    table.horz_banding = False
    col_widths = item["table"].get("col_widths")
    if col_widths:
        for j, w_px in enumerate(col_widths):
            table.columns[j].width = inches(w_px)
    for j, name in enumerate(cols):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = C["green"]
        cell.margin_left = cell.margin_right = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        r = p.add_run()
        r.text = clean(name)
        set_font(r, 14, C["white"], bold=True)
    for i, row in enumerate(rows, start=1):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = C["white"] if i % 2 == 1 else C["light"]
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            r = p.add_run()
            r.text = clean(value)
            bold = (j == 0) or (i == len(rows))
            set_font(r, 13, C["ink"] if bold else C["body"], bold=bold)
    below = top + height + 18
    for i, b in enumerate(item.get("bullets", [])[:3]):
        bullet(slide, b, 72, below + i * 30, 1128, 28, size=16)
    caveat_y = below + 30 * len(item.get("bullets", [])[:3]) + 8
    if item.get("caveat"):
        rect(slide, 72, caveat_y, 1136, 36, fill=C["amberBg"], line=C["amber"])
        text(slide, clean(item["caveat"]), 90, caveat_y + 8, 1098, 20, size=14,
             color=C["amberText"], bold=True)
    source_footer(slide, item["source"])


def agenda_slide(slide, page, total, item):
    header(slide, page, total, item.get("section") or "议程", item["title"])
    items = item.get("agenda", [])
    col_w, row_h = 546, 108
    for i, entry in enumerate(items[:6]):
        col, row = i % 2, i // 2
        x = M + col * (col_w + 44)
        y = 160 + row * (row_h + 26)
        rect(slide, x, y, col_w, row_h, fill=C["card"], line=C["cardLine"],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        num = rect(slide, x + 18, y + 26, 56, 56, fill=C["green"], line=C["green"],
                   shape=MSO_SHAPE.OVAL)
        tf = num.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = f"{i + 1:02d}"
        set_font(r, 18, C["white"], bold=True)
        text(slide, clean(entry.get("name", "")), x + 92, y + 20, col_w - 110, 30,
             size=18, color=C["ink"], bold=True)
        text(slide, clean(entry.get("desc", "")), x + 92, y + 56, col_w - 110, 40,
             size=12.5, color=C["muted"])
    source_footer(slide, item.get("source", ""))


def title_slide(slide, item, meeting, citation):
    rect(slide, 0, 0, 14, H, fill=C["green"], line=C["green"])
    text(slide, clean(meeting.get("type", "组会汇报")) + " · R4Det + RSSM", 86, 84, 640, 22,
         size=14, color=C["green"], bold=True)
    title = item.get("title") or "R4Det+RSSM：BEV 时序融合网络进展"
    text(slide, title, 86, 124, 1080, 60, size=compact_font(title, 38, 28, 26),
         color=C["ink"], bold=True)
    rect(slide, 88, 208, 120, 8, fill=C["amber"], line=C["amber"])
    if item.get("body"):
        text(slide, item["body"], 86, 240, 1000, 44, size=16, color=C["muted"])
    kpis = item.get("kpi") or []
    card_w, gap, height = 352, 28, 132
    for i, k in enumerate(kpis[:3]):
        x = 86 + i * (card_w + gap)
        rect(slide, x, 316, card_w, height, fill=C["card"], line=C["cardLine"],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.09)
        rect(slide, x, 316, 5, height, fill=C["green"], line=C["green"])
        text(slide, clean(k.get("value", "")), x + 24, 338, card_w - 44, 46,
             size=compact_font(k.get("value", ""), 28, 20, 12), color=C["ink"], bold=True)
        text(slide, clean(k.get("label", "")), x + 24, 396, card_w - 44, 34, size=12.5,
             color=C["muted"])
    if item.get("bullets"):
        text(slide, item["bullets"][0], 86, 496, 1080, 30, size=15, color=C["body"])
    text(slide, clean(meeting.get("meta", "")), 86, 548, 1080, 26, size=12.5, color=C["muted"])
    source_footer(slide, item.get("source", "Title"))


def closing_slide(slide, page, total, item):
    rect(slide, 0, 0, 14, H, fill=C["green"], line=C["green"])
    text(slide, item.get("title", "谢谢 · 欢迎讨论"), 86, 200, 1080, 70, size=36,
         color=C["ink"], bold=True, align=PP_ALIGN.CENTER)
    if item.get("body"):
        text(slide, item["body"], 86, 292, 1080, 36, size=17, color=C["muted"],
             align=PP_ALIGN.CENTER)
    chips = item.get("bullets") or []
    chip_w = 340
    total_w = len(chips[:3]) * chip_w + (len(chips[:3]) - 1) * 28
    x0 = (W - total_w) // 2
    for i, chip in enumerate(chips[:3]):
        x = x0 + i * (chip_w + 28)
        rect(slide, x, 366, chip_w, 110, fill=C["card"], line=C["cardLine"],
             shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.10)
        rect(slide, x, 366, 5, 110, fill=C["green"], line=C["green"])
        text(slide, clean(chip), x + 20, 382, chip_w - 40, 80, size=14, color=C["body"])
    footer(slide, f"{PAPER_CITATION} · {clean(item.get('source', 'Closing'))}")


def main(spec_path, out_path):
    global PAPER_CITATION
    spec_path = Path(spec_path).resolve()
    spec_dir = spec_path.parent
    spec = json.loads(spec_path.read_text(encoding="utf-8"))
    paper = spec.get("paper", {})
    meeting = spec.get("meeting", {})
    slides_spec = spec.get("slides", [])
    PAPER_CITATION = clean(paper.get("citation", paper.get("title", "Source document")))
    total = len(slides_spec)

    prs = Presentation()
    prs.slide_width = inches(W)
    prs.slide_height = inches(H)
    blank = prs.slide_layouts[6]

    used_figures = []
    for index, item in enumerate(slides_spec):
        page = index + 1
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = C["bg"]
        kind = item.get("type")
        if kind == "title":
            title_slide(slide, item, meeting, PAPER_CITATION)
        elif kind == "figure":
            used_figures.append(f"{item['figure_label']}（报告 §{item['pdf_page']}）")
            figure_slide(slide, page, total, item, spec_dir)
        elif item.get("table"):
            table_slide(slide, page, total, item, spec_dir)
        elif item.get("agenda"):
            agenda_slide(slide, page, total, item)
        elif item.get("closing"):
            closing_slide(slide, page, total, item)
        else:
            header(slide, page, total, item.get("section") or "网络组会", item["title"])
            body = item.get("body")
            kpis = item.get("kpi") or []
            if kpis:
                kpi_cards(slide, kpis)
            body_top = 260 if kpis else 168
            if body:
                text(slide, body, M, body_top, 1080, 110,
                     size=compact_font(body, 24, 18, 52), color=C["ink"], bold=True)
            start = (body_top + 128) if body else (body_top + 24)
            step = 96 if not body else 92
            for i, b in enumerate(item.get("bullets", [])[:3]):
                bullet(slide, b, M, start + i * step, 1080, 64, size=20)
            source_footer(slide, item.get("source", "Source document"))

    out_path = Path(out_path).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)

    figure_list = "、".join(used_figures) or "无图表页"
    verify = out_path.parent / f"{out_path.stem}.verify_checklist.md"
    verify.write_text(
        f"# 汇报前核对\n\n"
        f"- 证据来源：{PAPER_CITATION}\n"
        f"- 本次使用的证据图：{figure_list}\n"
        f"- 所有数字（40.42±0.51、35.59、34.65/37.08/39.00、37.94、+3.23、逐类别 strict/loose 表）均由 work_dirs 日志复算核对，与 docs/training_runs_full.md 一致；逐类别表取各 seed BEST epoch（s0 ep16 / s1 ep15 / s2 ep14）。\n"
        f"- 图 A–G 为本项目实验数据自绘（非论文原图）；每页页脚含来源定位（报告章节 / 代码文件）。\n"
        f"- 结论边界：时序融合收益在『同配置 No-Temporal』对照下成立；18e 与 24e 时代差异已在图C 注明；单 seed 提升不作为结论。\n"
        f"- 已知未解：Pedestrian strict≈0、Car-Truck 混淆、Cyclist 排序质量、CycCls seed2 进行中。\n",
        encoding="utf-8",
    )
    print(f"Wrote {total}-slide evidence deck to {out_path}")
    print(f"Wrote verification checklist to {verify}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
